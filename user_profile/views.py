from rest_framework import generics, filters
from django.contrib.auth.models import User
from .models import UserProfile, UserSkill, User, Skill
from .serializers import UserProfileSerializer, UserSkillSerializer
from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
import csv
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO
from docx import Document


class UserProfileList(generics.ListAPIView):
    queryset = UserProfile.objects.all()
    serializer_class = UserProfileSerializer
    filter_backends = [filters.SearchFilter]
    search_fields = ['skills__skill__name', 'skills__level']

class UserProfileDetail(generics.RetrieveUpdateDestroyAPIView):
    queryset = UserProfile.objects.all()
    serializer_class = UserProfileSerializer

class UserSkillCreate(generics.CreateAPIView):
    queryset = UserSkill.objects.all()
    serializer_class = UserSkillSerializer

class UserSkillUpdateDelete(generics.RetrieveUpdateDestroyAPIView):
    queryset = UserSkill.objects.all()
    serializer_class = UserSkillSerializer

def profiles_page(request):
    profiles = UserProfile.objects.all()
    return render(request, 'user_profile/profiles.html', {'profiles': profiles})

def edit_competency_page(request, user_id):
    user_profile = get_object_or_404(UserProfile, user_id=user_id)
    user_skills = UserSkill.objects.filter(user_profile=user_profile)
    return render(request, 'user_profile/edit_competency.html', {'user_skills': user_skills})


def add_skill_view(request, user_id):
    user = get_object_or_404(User, id=user_id)
    if request.method == 'POST':
        skill_name = request.POST.get('skill')
        skill_level = request.POST.get('level')
        Skill.objects.create(user=user, name=skill_name, level=skill_level)
        print(f"Skill '{skill_name}' at level '{skill_level}' added for user '{user.username}'")
        return redirect('profiles')
    return render(request, 'user_profile/add_skill.html', {'user': user})

def delete_skill(request, skill_id):
    skill = get_object_or_404(Skill, id=skill_id)
    if request.method == 'POST':
        skill.delete()
        return redirect('profiles')  # Redirect to profiles page after deletion
    return render(request, 'user_profile/delete_skill.html', {'skill': skill})

def update_skill(request, skill_id):
    skill = get_object_or_404(Skill, id=skill_id)
    if request.method == 'POST':
        skill.name = request.POST.get('name')
        skill.level = request.POST.get('level')
        skill.save()
        return redirect('profiles')  # Redirect to profiles page after update
    return render(request, 'user_profile/update_skill.html', {'skill': skill})

@login_required
def profiles_view(request):
    skill_query = request.GET.get('skill', '')
    level_query = request.GET.get('level', '')

    if skill_query or level_query:
        user_skills = Skill.objects.all()
        if skill_query:
            user_skills = user_skills.filter(name__icontains=skill_query)
        if level_query:
            user_skills = user_skills.filter(level__icontains=level_query)
        user_ids = user_skills.values_list('user_id', flat=True).distinct()
        users = User.objects.filter(id__in=user_ids)
    else:
        users = User.objects.all()

    context = {
        'users': users,
        'skill_query': skill_query,
        'level_query': level_query,
        'logged_in_user_id': request.user.id,  
    }
    return render(request, 'user_profile/profiles.html', context)


# def export_report(request):
#     format_type = request.GET.get('format', 'csv')
#     users = User.objects.all()

#     if format_type == 'pdf':
#         return generate_pdf_report(users)
#     elif format_type == 'doc':
#         return generate_docx_report(users)
#     else:
#         return generate_csv_report(users)

# def generate_csv_report(users):
#     response = HttpResponse(content_type='text/csv')
#     response['Content-Disposition'] = 'attachment; filename="employees.csv"'

#     writer = csv.writer(response)
#     writer.writerow(['First Name', 'Last Name', 'Skills', 'Levels'])

#     for user in users:
#         skills = ", ".join([f"{skill.name} ({skill.level})" for skill in user.skills.all()])
#         writer.writerow([user.first_name, user.last_name, skills])

#     return response

# def generate_pdf_report(users):
#     response = HttpResponse(content_type='application/pdf')
#     response['Content-Disposition'] = 'attachment; filename="employees.pdf"'

#     buffer = BytesIO()
#     p = canvas.Canvas(buffer, pagesize=letter)
#     width, height = letter

#     p.drawString(100, height - 40, "Employee Report")

#     y = height - 80
#     for user in users:
#         p.drawString(30, y, f"{user.first_name} {user.last_name}")
#         y -= 20
#         for skill in user.skills.all():
#             p.drawString(50, y, f"{skill.name} - {skill.level}")
#             y -= 20
#         y -= 20

#     p.showPage()
#     p.save()

#     buffer.seek(0)
#     return HttpResponse(buffer, content_type='application/pdf')

# def generate_doc_report(users):
#     response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
#     response['Content-Disposition'] = 'attachment; filename="employees.docx"'

#     doc = Document()
#     doc.add_heading('Employee Report', 0)

#     for user in users:
#         doc.add_heading(f"{user.first_name} {user.last_name}", level=1)
#         for skill in user.skills.all():
#             doc.add_paragraph(f"{skill.name} - {skill.level}")

#     buffer = BytesIO()
#     doc.save(buffer)
#     buffer.seek(0)
#     response.write(buffer.getvalue())

#     return response



@login_required
def export_report(request):
    skill_query = request.GET.get('skill', '')
    level_query = request.GET.get('level', '')
    format_type = request.GET.get('format', 'csv')
    
    user_skills = Skill.objects.all()
    if skill_query:
        user_skills = user_skills.filter(name__icontains=skill_query)
    if level_query:
        user_skills = user_skills.filter(level__icontains=level_query)
    user_ids = user_skills.values_list('user_id', flat=True).distinct()
    users = User.objects.filter(id__in=user_ids)

    if format_type == 'pdf':
        return generate_pdf_report(users)
    elif format_type == 'doc':
        return generate_docx_report(users)
    else:
        return generate_csv_report(users)


def generate_report(request):
    skill_query = request.GET.get('skill', '')
    level_query = request.GET.get('level', '')
    report_format = request.GET.get('format', 'csv')  # Default to CSV if not specified

    user_skills = Skill.objects.all()
    if skill_query:
        user_skills = user_skills.filter(name__icontains=skill_query)
    if level_query:
        user_skills = user_skills.filter(level__icontains=level_query)
    user_ids = user_skills.values_list('user_id', flat=True).distinct()
    users = User.objects.filter(id__in=user_ids)

    if report_format == 'csv':
        return generate_csv_report(users)
    elif report_format == 'pdf':
        return generate_pdf_report(users)
    elif report_format == 'docx':
        return generate_docx_report(users)
    else:
        return HttpResponse("Invalid format", status=400)

def generate_csv_report(users):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="report.csv"'
    writer = csv.writer(response)
    writer.writerow(['First Name', 'Last Name', 'Skills'])
    for user in users:
        skills = ', '.join([f"{skill.name} - {skill.level}" for skill in user.skills.all()])
        writer.writerow([user.first_name, user.last_name, skills])
    return response

def generate_pdf_report(users):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setTitle("User Report")
    pdf.drawString(100, 750, "User Report")
    y = 700
    for user in users:
        pdf.drawString(100, y, f"Name: {user.first_name} {user.last_name}")
        y -= 20
        for skill in user.skills.all():
            pdf.drawString(120, y, f"{skill.name} - {skill.level}")
            y -= 20
        y -= 20
    pdf.save()
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="report.pdf"'
    return response

def generate_docx_report(users):
    doc = Document()
    doc.add_heading('User Report', 0)
    for user in users:
        doc.add_heading(f"{user.first_name} {user.last_name}", level=1)
        for skill in user.skills.all():
            doc.add_paragraph(f"{skill.name} - {skill.level}")
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="report.docx"'
    doc.save(response)
    return response
